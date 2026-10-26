#!/usr/bin/env python3
"""Build report diagrams locally using Graphviz and Matplotlib; no network required."""
from __future__ import annotations
import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import textwrap
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "qtmnc-mpl-cache"))
os.environ.setdefault("XDG_CACHE_HOME", str(Path(tempfile.gettempdir()) / "qtmnc-font-cache"))


def fonts():
    pairs = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ]
    for regular, bold in pairs:
        if Path(regular).is_file() and Path(bold).is_file():
            return regular, bold
    raise RuntimeError("Cần font Noto Sans, DejaVu Sans hoặc Arial có dấu tiếng Việt")


def draw_graph(item):
    stem = item["stem"]
    if stem.startswith("06-"):
        return
    source = ROOT / "source" / (stem + ".dot")
    for fmt in ("svg", "png"):
        output = ROOT / fmt / (stem + "." + fmt)
        command = ["dot", str(source), "-T" + fmt, "-o", str(output)]
        if fmt == "png":
            command.append("-Gdpi=220")
        result = subprocess.run(command, text=True, capture_output=True)
        if result.returncode:
            raise RuntimeError(f"Graphviz lỗi ở {stem}: {result.stderr}")


def draw_sequence():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import FancyBboxPatch
    regular, bold = fonts()
    font_manager.fontManager.addfont(regular)
    font_manager.fontManager.addfont(bold)
    family = font_manager.FontProperties(fname=regular).get_name()
    plt.rcParams.update({"font.family": family, "svg.fonttype": "none", "axes.unicode_minus": False})
    fig, ax = plt.subplots(figsize=(10, 12))
    fig.subplots_adjust(left=.025, right=.975, top=.975, bottom=.025)
    ax.set(xlim=(0, 1), ylim=(0, 1))
    ax.axis("off")
    positions = {"M": .1, "L": .3, "D": .5, "W": .7, "T": .9}
    titles = {"M": "Windows\nLuồng đo", "L": "Linux\nSSH Server", "D": "SQLite", "W": "Luồng gửi", "T": "Telegram\nBot API"}
    for key, x in positions.items():
        ax.plot([x, x], [.025, .92], color="#94a3b8", linewidth=1, linestyle=(0, (4, 4)), zorder=1)
        ax.add_patch(FancyBboxPatch((x-.087, .927), .174, .053, boxstyle="round,pad=0.005",
                                  linewidth=1.4, edgecolor="#2563eb", facecolor="#eaf2ff", zorder=3))
        ax.text(x, .9535, titles[key], ha="center", va="center", fontsize=12.5, color="#0f172a", fontweight="bold")
    events = [
        ("M", "L", "Đọc CPU / RAM / interface", "request"),
        ("L", "M", "RAM = 70%", "response"),
        (None, None, "Ngưỡng RAM 65% • Trạng thái trước: OK", "note"),
        ("M", "D", "Transaction: RAM → ALERT + tin cảnh báo", "request"),
        ("W", "D", "Lấy tin cũ nhất đến hạn", "request"),
        ("D", "W", "Tin cảnh báo RAM", "response"),
        ("W", "T", "POST lần 1: lỗi kết nối", "failed"),
        ("W", "D", "Giữ tin, tăng attempts, hẹn thử lại", "request"),
        (None, None, "Vòng đo tiếp tục khi luồng gửi đang chờ thử lại", "note"),
        ("W", "D", "Đọc lại tin đến hạn", "request"),
        ("W", "T", "POST lần 2", "request"),
        ("T", "W", "HTTP 200, ok=true", "response"),
        ("W", "D", "Xóa tin đã xác nhận", "request"),
        ("M", "L", "Chu kỳ sau: đọc RAM", "request"),
        ("L", "M", "RAM = 50%", "response"),
        (None, None, "50% ≤ 60% → đủ điều kiện phục hồi", "note"),
        ("M", "D", "Transaction: RAM → OK + tin phục hồi", "request"),
        ("W", "D", "Lấy tin phục hồi đến hạn", "request"),
        ("W", "T", "Gửi tin RAM bình thường", "request"),
        ("T", "W", "HTTP 200, ok=true", "response"),
        ("W", "D", "Xóa tin đã xác nhận", "request"),
    ]
    y, step, number = .88, .0395, 0
    for source, target, text, kind in events:
        if kind == "note":
            ax.add_patch(FancyBboxPatch((.04, y-.007), .92, .03, boxstyle="round,pad=0.003",
                                      linewidth=.8, edgecolor="#d97706", facecolor="#fffbeb", zorder=4))
            ax.text(.5, y+.008, text, ha="center", va="center", fontsize=11.5, color="#78350f", zorder=5)
        else:
            number += 1
            x1, x2 = positions[source], positions[target]
            color = "#e11d48" if kind == "failed" else "#2563eb" if kind == "request" else "#059669"
            style = "--" if kind in {"response", "failed"} else "-"
            ax.annotate("", xy=(x2, y), xytext=(x1, y),
                        arrowprops={"arrowstyle": "->", "color": color, "lw": 1.4, "linestyle": style}, zorder=4)
            if kind == "failed":
                ax.plot(x2, y, marker="x", color=color, markersize=8, markeredgewidth=2, zorder=6)
            width = 23 if abs(x2-x1) < .3 else 43
            label = textwrap.fill(f"{number:02d}. {text}", width=width)
            ax.text((x1+x2)/2, y+.0045, label, fontsize=11.5, ha="center", va="bottom", color="#0f172a",
                    linespacing=1.05, bbox={"facecolor":"white","edgecolor":"none","pad":1}, zorder=5)
        y -= step
    stem = "06-trinh-tu-ram-telegram"
    fig.savefig(ROOT / "png" / (stem+".png"), dpi=240, facecolor="white")
    fig.savefig(ROOT / "svg" / (stem+".svg"), facecolor="white")
    plt.close(fig)


def report_pdf(manifest):
    from PIL import Image
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    regular, bold = fonts()
    pdfmetrics.registerFont(TTFont("ReportSans", regular))
    pdfmetrics.registerFont(TTFont("ReportSansBold", bold))
    pdfmetrics.registerFontFamily("ReportSans", normal="ReportSans", bold="ReportSansBold")
    path = ROOT / "Bo_so_do_giam_sat_thiet_bi.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Bộ sơ đồ — Hệ thống giám sát thiết bị và cảnh báo Telegram/Slack")
    c.setAuthor("qtmnc")
    title = ParagraphStyle("Title", fontName="ReportSansBold", fontSize=22, leading=30, textColor=HexColor("#0f172a"))
    heading = ParagraphStyle("Heading", fontName="ReportSansBold", fontSize=18, leading=25, textColor=HexColor("#0f172a"))
    body = ParagraphStyle("Body", fontName="ReportSans", fontSize=10.5, leading=16, textColor=HexColor("#334155"))
    def paragraph(text, x, y, width, style=body):
        p = Paragraph(escape(text).replace("\n", "<br/>"), style)
        _, height = p.wrap(width, 2000)
        p.drawOn(c, x, y-height)
        return y-height
    def footer(page, label):
        c.setFont("ReportSans", 8)
        c.setFillColor(HexColor("#64748b"))
        width, _ = c._pagesize
        c.drawString(36, 22, label)
        c.drawRightString(width-36, 22, str(page))
    width, height = A4
    y = height-100
    y = paragraph("HỆ THỐNG GIÁM SÁT\nTRẠNG THÁI THIẾT BỊ", 42, y, width-84, title)-20
    y = paragraph("Và gửi cảnh báo qua Telegram / Slack",42,y,width-84,heading)-35
    y = paragraph("BỘ 8 MÔ HÌNH VÀ SƠ ĐỒ THIẾT KẾ",42,y,width-84,heading)-20
    y = paragraph("Mô hình thực hành: máy thật Windows chạy Python, giám sát máy ảo Linux qua SSH và gửi cảnh báo Telegram.",42,y,width-84)-25
    y = paragraph("Có ảnh PNG, vector SVG, mã Mermaid và mã nguồn dựng hình. Các kịch bản số đo trong sơ đồ trình tự là minh họa thiết kế.",42,y,width-84)-25
    paragraph("Kho mã: https://github.com/tunato2807-del/qtmnc\nChương trình dùng dashboard Terminal, không phải dashboard web.",42,y,width-84)
    footer(1,"Tài liệu sơ đồ thiết kế • Windows → Linux VM")
    c.showPage()
    y = paragraph("Phạm vi và cấu hình mô hình",36,height-40,width-72,heading)-20
    rows = [
        ["Hạng mục","Cấu hình / ý nghĩa"],
        ["Máy giám sát","Windows • Python • VS Code"],
        ["Thiết bị đích","Linux VM • 192.168.73.129/24 • ens33"],
        ["Windows VMnet8","192.168.73.1/24"],
        ["Gateway NAT của VM","192.168.73.2"],
        ["Thu thập chính","SSH / TCP 22, kiểm tra host key"],
        ["Chu kỳ","30 giây mặc định • gần thời gian thực"],
        ["Kênh đang dùng","Telegram Bot API qua HTTPS / TCP 443"],
        ["Tùy chọn trong mã","SNMP / UDP 161 và Slack Webhook"],
        ["CPU / RAM","Cảnh báo >75% / >65%; phục hồi ≤70% / ≤60%"],
        ["Nhiệt độ","Đang tắt trong demo VM; không giả thành 0°C"],
        ["Lưu trữ","SQLite: trạng thái gần nhất và tin chưa gửi"],
    ]
    table_rows = [[Paragraph(escape(v),body) for v in row] for row in rows]
    table = Table(table_rows, colWidths=[145,width-72-145], hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),HexColor("#dbeafe")),("GRID",(0,0),(-1,-1),.5,HexColor("#cbd5e1")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    _, table_height = table.wrap(width-72,700)
    table.drawOn(c,36,y-table_height)
    y -= table_height+20
    for text in [
        "Nét liền: luồng của bản demo chính. Nét đứt vàng: thành phần tùy chọn. Màu xanh lá: lưu trữ hoặc kết quả hợp lệ; màu vàng: quyết định, thiếu dữ liệu hoặc thử lại.",
        "VM nằm trên cùng host Windows: khi host tắt/ngủ, VM có thể dừng. Muốn phát hiện host tắt hoàn toàn cần máy giám sát độc lập.",
        "Nguồn đối chiếu: monitor_ssh.py, terminal_ui.py, cấu hình và log được cung cấp. Tài liệu không khẳng định đã thử SNMP/Slack/nhiệt độ trên thiết bị của người dùng.",
    ]:
        y = paragraph(text,36,y,width-72)-12
    footer(2,"Phạm vi và chú giải")
    c.showPage()
    for i,item in enumerate(manifest,1):
        image = ROOT / "png" / (item["stem"]+".png")
        with Image.open(image) as picture:
            iw,ih=picture.size
        page = landscape(A4) if iw/ih > 1.25 else A4
        c.setPageSize(page)
        width,height = page
        y = paragraph(f"Hình {i}. {item['title']}",36,height-30,width-72,heading)-12
        y = paragraph(item["caption"],36,y,width-72)-10
        note_height = Paragraph(escape(item["notes"]),body).wrap(width-72,2000)[1]
        available_height = y-(55+note_height+15)
        scale=min((width-72)/iw,available_height/ih)
        dw,dh=iw*scale,ih*scale
        x=(width-dw)/2
        c.drawImage(str(image),x,y-dh,width=dw,height=dh,mask="auto")
        paragraph(item["notes"],36,y-dh-12,width-72)
        footer(i+2, "qtmnc • " + item["chapter"])
        c.showPage()
    c.save()
    return path


def report_docx(manifest):
    from PIL import Image
    from docx import Document
    from docx.enum.section import WD_SECTION_START, WD_ORIENT
    from docx.shared import Inches, Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    normal=doc.styles["Normal"]
    normal.font.name="Arial"
    normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.1
    for name in ("Title","Heading 1","Heading 2"):
        doc.styles[name].font.name="Arial"
    def size(section,wide=False):
        section.orientation=WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        section.page_width=Cm(29.7 if wide else 21)
        section.page_height=Cm(21 if wide else 29.7)
        section.top_margin=section.bottom_margin=Cm(1.25)
        section.left_margin=section.right_margin=Cm(1.25)
    size(doc.sections[0])
    doc.add_heading("HỆ THỐNG GIÁM SÁT\nTRẠNG THÁI THIẾT BỊ",0)
    doc.add_heading("Và gửi cảnh báo qua Telegram / Slack",1)
    doc.add_paragraph("BỘ 8 MÔ HÌNH VÀ SƠ ĐỒ THIẾT KẾ")
    doc.add_paragraph("Mô hình thực hành: Windows chạy Python → Linux VM qua SSH → Telegram.")
    doc.add_paragraph("Mã nguồn và bản dựng: https://github.com/tunato2807-del/qtmnc")
    doc.add_paragraph("Các hình có PNG để chèn nhanh, SVG để phóng to, Mermaid và mã nguồn để chỉnh sửa.")
    doc.add_page_break()
    doc.add_heading("Phạm vi mô hình",1)
    table=doc.add_table(rows=1,cols=2)
    table.style="Table Grid"
    table.rows[0].cells[0].text="Hạng mục"
    table.rows[0].cells[1].text="Giá trị"
    for key,value in [
        ("Máy giám sát","Windows / Python / VS Code"),
        ("Máy đích","Linux VM: 192.168.73.129, interface ens33"),
        ("VMnet8 / gateway VM","192.168.73.1 / 192.168.73.2"),
        ("Thu thập / chu kỳ","SSH TCP 22 / 30 giây"),
        ("Cảnh báo","Telegram HTTPS 443; Slack là tùy chọn"),
        ("Ngưỡng CPU","Cảnh báo >75%, phục hồi ≤70%"),
        ("Ngưỡng RAM","Cảnh báo >65%, phục hồi ≤60%"),
        ("Nhiệt độ","Đang tắt trong demo VM"),
        ("Tùy chọn","SNMP UDP 161; chưa xác minh trong demo"),
        ("Dữ liệu","Trạng thái gần nhất + outbox SQLite"),
    ]:
        row=table.add_row().cells
        row[0].text,row[1].text=key,value
    doc.add_paragraph("Nét đứt vàng biểu diễn phần tùy chọn. Kịch bản RAM trong hình trình tự là minh họa, không phải bản ghi một phép thử đã được xác nhận.")
    doc.add_paragraph("Nếu Windows host tắt hoặc ngủ, Linux VM có thể dừng. Muốn giám sát tính sẵn sàng của host cần máy giám sát độc lập.")
    for i,item in enumerate(manifest,1):
        image=ROOT/"png"/(item["stem"]+".png")
        with Image.open(image) as picture:
            iw,ih=picture.size
        section=doc.add_section(WD_SECTION_START.NEW_PAGE)
        size(section,iw/ih>1.25)
        doc.add_heading(f"Hình {i}. {item['title']}",1)
        doc.add_paragraph(item["caption"])
        width=(section.page_width-section.left_margin-section.right_margin)/914400
        maxheight=(section.page_height-section.top_margin-section.bottom_margin)/914400-2.25
        dw=min(width, maxheight*iw/ih)
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(4)
        p.add_run().add_picture(str(image),width=Inches(dw))
        doc.add_paragraph(item["notes"])
    footer=doc.sections[0].footer.paragraphs[0]
    footer.add_run("qtmnc • Tài liệu sơ đồ thiết kế • Trang ")
    field=OxmlElement("w:fldSimple")
    field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    path=ROOT/"Bo_so_do_giam_sat_thiet_bi.docx"
    doc.save(path)
    return path


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reports-only",action="store_true")
    args=parser.parse_args()
    manifest=json.loads((ROOT/"manifest.json").read_text(encoding="utf-8"))
    if not args.reports_only:
        if shutil.which("dot") is None:
            raise RuntimeError("Cần Graphviz và lệnh dot trong PATH")
        for fmt in ("png","svg"):
            (ROOT/fmt).mkdir(exist_ok=True)
        with ThreadPoolExecutor(max_workers=3) as pool:
            list(pool.map(draw_graph,manifest))
        draw_sequence()
        # Generated path coordinates may contain redundant trailing whitespace.
        # Normalize text SVG only; never rewrite PDF or other binary artifacts.
        for path in (ROOT/"svg").glob("*.svg"):
            path.write_text("\n".join(line.rstrip() for line in path.read_text(encoding="utf-8").splitlines())+"\n", encoding="utf-8")
    pdf=report_pdf(manifest)
    word=report_docx(manifest)
    print(f"Đã dựng {len(manifest)} sơ đồ PNG/SVG.")
    print(pdf)
    print(word)


if __name__=="__main__":
    main()
