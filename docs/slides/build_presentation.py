#!/usr/bin/env python3
"""Build the editable 22-slide Vietnamese QTMNC project presentation."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
W, H = 13.333, 7.5

NAVY = "14243A"
NAVY2 = "1B3552"
INK = "1E2D40"
MUTED = "607086"
TEAL = "15A995"
TEAL_DARK = "087C71"
BLUE = "4C78E8"
ORANGE = "E7A23B"
RED = "D65F5B"
GREEN = "268D68"
PURPLE = "8165C6"
BG = "F4F7FB"
WHITE = "FFFFFF"
LINE = "DFE6EF"
PALE_BLUE = "EAF0FF"
PALE_TEAL = "E7F6F3"
PALE_ORANGE = "FFF3DF"
PALE_RED = "FCE9E8"
PALE_GREEN = "E5F3EC"
PALE_PURPLE = "F0ECFA"
FONT = "Aptos"

SLIDES: list[dict[str, Any]] = [
    {
        "section": "TÓM TẮT ĐIỀU HÀNH", "title": "Giám sát thiết bị và cảnh báo",
        "message": "Windows chạy Python kiểm tra Ubuntu định kỳ; khi trạng thái đổi, hệ thống gửi cảnh báo Telegram/Slack.",
        "points": [
            ("Đối tượng", "Ubuntu VMware; theo dõi CPU, RAM, nhiệt độ và interface."),
            ("Đường chính", "SSH qua GNS3; bản đơn giản chờ 5 giây sau mỗi lượt đọc."),
            ("Mở rộng", "SNMP/hybrid; gửi cảnh báo qua HTTPS tới Telegram/Slack."),
        ],
        "evidence": "PHẠM VI  4 chỉ số  ·  SSH TCP/22  ·  polling gần thời gian thực",
        "visual": {"kind": "cover", "nodes": ["WINDOWS\nPython / VS Code", "GNS3\nMạng thử nghiệm", "UBUNTU VM\nThiết bị đích", "Telegram\nSlack"]},
        "source": "README.md · docs/ARCHITECTURE.md · examples/simple.env",
    },
    {
        "section": "BỐI CẢNH", "title": "Bối cảnh và mục tiêu hệ thống",
        "message": "Giải pháp biến số liệu thiết bị thành cảnh báo và bằng chứng xử lý sự cố có thể trình diễn.",
        "points": [
            ("Vấn đề", "Lỗi CPU, bộ nhớ, nhiệt độ hoặc interface có thể bị phát hiện muộn."),
            ("Mục tiêu", "Kiểm tra có lịch; phát hiện chuyển trạng thái; báo người quản trị."),
            ("Phạm vi", "Một Windows giám sát Ubuntu VM qua mạng GNS3 hiện có."),
        ],
        "evidence": "ĐẦU RA  Dashboard · cảnh báo lỗi/phục hồi · hàng đợi tin chưa gửi",
        "visual": {"kind": "flow", "nodes": ["THIẾT BỊ\nTài nguyên / mạng", "PYTHON\nĐọc và đánh giá", "NGƯỜI QUẢN TRỊ\nNhận thông báo"], "colors": [BLUE, ORANGE, TEAL]},
        "source": "README.md · docs/ARCHITECTURE.md",
    },
    {
        "section": "KIẾN TRÚC", "title": "Kiến trúc tổng thể giải pháp",
        "message": "Thu thập, đánh giá, lưu trữ và gửi tin là các bước tách biệt trong cùng một ứng dụng.",
        "points": [
            ("Máy giám sát", "Windows chạy monitor.py và bảng Rich trong Terminal."),
            ("Thiết bị đích", "Ubuntu trả số liệu qua SSH; SNMP là collector mở rộng."),
            ("Kênh gửi", "Luồng gửi riêng chuyển tin qua Telegram Bot API/Slack Webhook."),
        ],
        "evidence": "MÔ HÌNH  Python Windows → GNS3 → Ubuntu; HTTPS → dịch vụ nhắn tin",
        "visual": {"kind": "architecture"},
        "source": "docs/ARCHITECTURE.md · docs/truc_quan/README.md",
    },
    {
        "section": "THÀNH PHẦN", "title": "Thiết bị và công cụ thực hành",
        "message": "Mỗi thiết bị đảm nhiệm một vai trò; chỉ Ubuntu là đối tượng lấy CPU/RAM trong lab hiện tại.",
        "points": [
            ("Windows", "Máy thật; chạy VS Code, Python, cấu hình và gửi cảnh báo."),
            ("VMware", "Cấp CPU/RAM ảo, Ubuntu và các card mạng cho thiết bị đích."),
            ("GNS3", "Hai Cloud và Switch1 chuyển tiếp mạng; không tự đo tài nguyên."),
        ],
        "evidence": "ĐỐI TƯỢNG ĐO  Số CPU/RAM lấy từ Ubuntu VM, không phải toàn Windows",
        "visual": {"kind": "roles"},
        "source": "docs/GNS3_LAB.md · docs/truc_quan/README.md",
    },
    {
        "section": "MẠNG", "title": "Mô hình mạng hiện tại trên GNS3",
        "message": "Đường giám sát dùng subnet 192.168.10.0/24; ens37 quản trị và ens33 là interface theo dõi.",
        "points": [
            ("Windows", "Card TAP Local Area Connection: 192.168.10.1/24."),
            ("Đường SSH", "Cloud-WINDOWS → Switch1 → Cloud-LINUX/VMnet2 → ens37."),
            ("Interface đo", "ens33 NAT; đã ghi nhận 192.168.73.129/24, gateway .73.2."),
        ],
        "evidence": "SSH ĐÍCH  192.168.10.129:22  ·  ens37 192.168.10.129/24",
        "visual": {"kind": "network"},
        "source": "WINDOWS_TAP_GNS3.md · docs/truc_quan/README.md · kết quả người dùng",
    },
    {
        "section": "MẠNG", "title": "Đường SSH và vai trò TAP",
        "message": "TAP/UDP chỉ nối Windows vào lab GNS3 hiện tại; SSH mới vận chuyển yêu cầu lấy số liệu.",
        "points": [
            ("SSH", "Phiên TCP/22 được thiết lập từ Windows tới Ubuntu .10.129."),
            ("Cầu nối", "TAP-Windows + UDP localhost đưa Ethernet vào Cloud GNS3."),
            ("Phạm vi", "TAP không đo CPU/RAM, không đánh giá và không gửi cảnh báo."),
        ],
        "evidence": "UDP CẦU NỐI  Cloud local 30010 ↔ helper 127.0.0.1:30011",
        "visual": {"kind": "tappath"},
        "source": "WINDOWS_TAP_GNS3.md · tools/windows_tap_bridge.py",
    },
    {
        "section": "THU THẬP", "title": "Lựa chọn giao thức thu thập",
        "message": "SSH, SNMP và hybrid dùng các cách lấy số liệu khác nhau nhưng chung phần đánh giá và cảnh báo.",
        "points": [
            ("SSH", "TCP/22; chạy lệnh Linux qua Paramiko; là bản đơn giản hiện dùng."),
            ("SNMP", "UDP/161; hỏi agent bằng OID; cần agent và cấu hình tương ứng."),
            ("Hybrid", "Ưu tiên SSH cho metric; ưu tiên SNMP cho interface; có dự phòng."),
        ],
        "evidence": "HYBRID  Mất từng dịch vụ vẫn hiển thị trạng thái; nguồn còn lại có thể tiếp tục",
        "visual": {"kind": "protocols"},
        "source": "device_monitor/collectors/{ssh,snmp,hybrid}.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "PHẦN MỀM", "title": "Các module phần mềm chính",
        "message": "Các module tách thu thập, trạng thái và gửi cảnh báo để dễ giải thích và mở rộng.",
        "points": [
            ("Điểm vào", "monitor.py điều phối cấu hình, SSH, bảng và vòng lặp."),
            ("Xử lý", "Collectors lấy số liệu; engine quyết định OK/ALERT/UNKNOWN."),
            ("Đầu ra", "storage.py lưu SQLite; notifications.py gửi và retry."),
        ],
        "evidence": "MÃ CHÍNH  config → collectors → engine → storage → notifications → ui",
        "visual": {"kind": "modules"},
        "source": "monitor.py · device_monitor/app.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "CHỈ SỐ", "title": "Bốn chỉ số thiết bị cần theo dõi",
        "message": "Bản đơn giản thu thập bốn nhóm chỉ số trên Ubuntu; một chỉ số thiếu không phủ định các chỉ số khác.",
        "points": [
            ("CPU", "Hai lần đọc /proc/stat; chênh lệch biểu thị thời gian bận."),
            ("RAM", "/proc/meminfo; dùng MemTotal và MemAvailable."),
            ("Nhiệt độ/interface", "Cảm biến /sys nếu có; operstate của ens33."),
        ],
        "evidence": "NGUỒN  CPU/RAM/temp/interface đều thuộc Ubuntu VM",
        "visual": {"kind": "metrics"},
        "source": "device_monitor/collectors/ssh.py · examples/simple.env",
    },
    {
        "section": "CHỈ SỐ", "title": "Cách hệ thống tính CPU và RAM",
        "message": "Giá trị được chuẩn hóa thành phần trăm từ bộ đếm Linux và lượng RAM còn dùng được.",
        "points": [
            ("CPU", "100 × (1 − idle delta / total delta); đọc cách nhau khoảng 1 giây."),
            ("RAM", "100 × (MemTotal − MemAvailable) / MemTotal."),
            ("Diễn giải", "4 GB cấp cho VM, còn 1 GB khả dụng → RAM dùng xấp xỉ 75%."),
        ],
        "evidence": "VÍ DỤ  3/4 GB sử dụng = 75% RAM Ubuntu",
        "visual": {"kind": "formulas"},
        "source": "device_monitor/collectors/ssh.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "CHỈ SỐ", "title": "Nhiệt độ và tính toàn vẹn số liệu",
        "message": "Không có cảm biến phải được báo UNKNOWN hoặc bỏ theo cấu hình, không đổi thành nhiệt độ 0°C.",
        "points": [
            ("Auto", "Đọc cảm biến CPU đã nhận diện trong /sys/hwmon hoặc thermal."),
            ("UNKNOWN", "Máy ảo thường không lộ cảm biến vật lý; ghi nhận thiếu dữ liệu."),
            ("Demo", "Đọc mẫu có hạn dùng; mọi số đo/cảnh báo phải gắn nhãn MÔ PHỎNG."),
        ],
        "evidence": "PROFILE  MONITOR_TEMP_ENABLED · TEMP_SOURCE=auto/demo",
        "visual": {"kind": "temperature"},
        "source": "device_monitor/probe.py · lab/linux/labctl.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "VẬN HÀNH", "title": "Chu kỳ thu thập và độ trễ",
        "message": "Polling 5 giây là gần thời gian thực; phát hiện và giao tin vẫn chịu độ trễ đo, mạng và API.",
        "points": [
            ("Lấy mẫu", "Kết nối, đọc chỉ số, đánh giá và cập nhật bảng mỗi vòng."),
            ("Chu kỳ", "Profile đơn giản chờ 5 giây sau lượt đọc; CPU cần mẫu ~1 giây."),
            ("Độ trễ", "Sự cố chờ lượt kế tiếp; timeout/retry/API làm tin tới muộn."),
        ],
        "evidence": "5 GIÂY  Không phải cam kết mọi cảnh báo đến trong đúng 5 giây",
        "visual": {"kind": "cycle"},
        "source": "examples/simple.env · docs/ARCHITECTURE.md · docs/TEST_PLAN.md",
    },
    {
        "section": "CẢNH BÁO", "title": "Ngưỡng cảnh báo và phục hồi",
        "message": "Ngưỡng phục hồi có khoảng đệm 5 điểm để giảm đổi trạng thái liên tục.",
        "points": [
            ("CPU", "Cảnh báo >75%; sau đó phục hồi khi ≤70%."),
            ("RAM", "Cảnh báo >65%; sau đó phục hồi khi ≤60%."),
            ("Nhiệt độ", "Cảnh báo >65°C; sau đó phục hồi khi ≤60°C."),
        ],
        "evidence": "HYSTERESIS  RECOVERY_MARGIN=5; interface DOWN→ALERT, UP→OK",
        "visual": {"kind": "thresholds"},
        "source": "examples/simple.env · device_monitor/engine.py",
    },
    {
        "section": "CẢNH BÁO", "title": "Máy trạng thái và hysteresis",
        "message": "Chỉ tạo tin khi trạng thái chuyển; UNKNOWN giữ trạng thái hợp lệ gần nhất để xử lý phục hồi đúng.",
        "points": [
            ("OK", "Giá trị trong miền bình thường; không gửi lại mỗi vòng."),
            ("ALERT", "Vượt ngưỡng; giữ cảnh báo tới khi đạt ngưỡng phục hồi."),
            ("UNKNOWN", "Thiếu dữ liệu; không suy diễn là bình thường hoặc bằng 0."),
        ],
        "evidence": "RAM  50 → 72 → 74 → 63 → 58: một cảnh báo, một tin phục hồi",
        "visual": {"kind": "states"},
        "source": "device_monitor/engine.py · device_monitor/storage.py",
    },
    {
        "section": "CẢNH BÁO", "title": "Phân biệt ba nhóm sự cố",
        "message": "Ngắt interface đo, mất đường SSH và mất Internet gây ba kết quả vận hành khác nhau.",
        "points": [
            ("ens33 DOWN", "SSH qua ens37 còn; đọc trạng thái interface và báo lỗi."),
            ("Mất SSH", "Không có số đo mới; đánh dấu mất kết nối và báo phục hồi."),
            ("Mất Internet", "Lấy mẫu vẫn chạy; cảnh báo xếp hàng chờ gửi."),
        ],
        "evidence": "NGUYÊN TẮC  Ngắt ens33 để thử trạng thái; ngắt GNS3 để thử mất SSH",
        "visual": {"kind": "incidents"},
        "source": "docs/TEST_PLAN.md · docs/truc_quan/README.md",
    },
    {
        "section": "GỬI TIN", "title": "Hàng đợi và độ tin cậy cảnh báo",
        "message": "SQLite lưu trạng thái và tin trong cùng giao dịch; luồng gửi riêng thử lại các tin chưa xác nhận.",
        "points": [
            ("Outbox", "Giữ tin chưa giao qua lần dừng/chạy lại ứng dụng."),
            ("Retry", "Backoff 2, 4, 8… tối đa 300 giây; HTTP 429 đọc retry-after."),
            ("Cam kết", "At-least-once; phản hồi thất lạc có thể làm tin bị gửi trùng."),
        ],
        "evidence": "SQLite  states + outbox; mỗi kênh giữ thứ tự tin riêng",
        "visual": {"kind": "outbox"},
        "source": "device_monitor/storage.py · device_monitor/notifications.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "AN TOÀN", "title": "Bảo vệ thông tin và kết nối",
        "message": "Xác minh host key, lưu secret cục bộ và giới hạn giao thức quản trị trong mạng lab.",
        "points": [
            ("SSH", "RejectPolicy + known_hosts; không tự tin host key lạ."),
            ("Secret", "Mật khẩu/token/webhook ở .env bị Git ignore; không ghi log."),
            ("SNMP", "v2c chỉ lab cô lập; v3 authPriv hỗ trợ xác thực/mã hóa."),
        ],
        "evidence": "PORT  SSH TCP/22 · SNMP UDP/161 · API HTTPS/TCP 443",
        "visual": {"kind": "security"},
        "source": "device_monitor/collectors/ssh.py · device_monitor/config.py · docs/ARCHITECTURE.md",
    },
    {
        "section": "TRIỂN KHAI", "title": "Hai profile triển khai dự án",
        "message": "Bản đơn giản giảm bước cài đặt; profile GNS3 mở rộng SNMP/hybrid và mạng thử nghiệm riêng.",
        "points": [
            ("Đơn giản", "run_windows.cmd; SSH tới .10.129; theo dõi ens33."),
            ("Mở rộng", ".env.gns3; SSH/SNMP/hybrid; gns3test .30.129."),
            ("Dùng chung", "Ngưỡng, SQLite, dashboard và Telegram/Slack."),
        ],
        "evidence": "LAB MỞ RỘNG  VMnet3 .30.0/24 · VPCS .30.10 (kiểm tra mạng)",
        "visual": {"kind": "profiles"},
        "source": "README.md · docs/GNS3_LAB.md · examples/{simple,gns3}.env",
    },
    {
        "section": "NGHIỆM THU", "title": "Kế hoạch kiểm thử và nghiệm thu",
        "message": "Mỗi tình huống phải gắn kích thích, phản ứng mong đợi và bằng chứng trên máy thật.",
        "points": [
            ("Tài nguyên", "Tạo tải CPU/RAM; xác nhận cảnh báo và phục hồi."),
            ("Kết nối", "Ngắt ens33, SSH/GNS3 hoặc Internet riêng theo từng ca."),
            ("Dữ liệu", "Kiểm tra nhiệt độ UNKNOWN/demo và host key sai."),
        ],
        "evidence": "BẰNG CHỨNG  Thời điểm đo · tin nhận · ảnh đã che secret · phiên bản/commit",
        "visual": {"kind": "testmatrix"},
        "source": "docs/TEST_PLAN.md · docs/GNS3_LAB.md",
    },
    {
        "section": "KIỂM CHỨNG", "title": "Kết quả kiểm chứng hiện có",
        "message": "Kiểm thử mã bao phủ collector, ngưỡng và retry; xác nhận trên Windows/GNS3 thật vẫn là bước riêng.",
        "points": [
            ("Tự động", "119 unit tests; SSH/SNMP loopback và parser/test outbox."),
            ("Mạng thật", "TCP/22 tới .10.129 từng thành công qua Local Area Connection."),
            ("Giới hạn chứng cứ", "Loopback tests không chứng minh Telegram/SNMP agent thật."),
        ],
        "evidence": "ẢNH NGƯỜI DÙNG  TAP counters: 106→GNS3, 57→TAP, dropped 0 (một thời điểm)",
        "visual": {"kind": "evidence"},
        "source": "docs/VALIDATION.md · ảnh/test người dùng · docs/TEST_PLAN.md",
    },
    {
        "section": "GIỚI HẠN", "title": "Giới hạn kỹ thuật cần công bố",
        "message": "Kết quả phải nêu giới hạn môi trường, phạm vi đo và mức xác nhận thực tế.",
        "points": [
            ("Nhiệt độ", "VM có thể thiếu cảm biến; demo không phải đo phần cứng."),
            ("Lịch sử", "SQLite lưu state/outbox, không phải kho thời gian-series."),
            ("Cùng máy", "Windows tắt/ngủ có thể làm monitor và VM cùng dừng."),
        ],
        "evidence": "PHẠM VI  Một thiết bị / tiến trình / state file theo cấu hình",
        "visual": {"kind": "limitations"},
        "source": "docs/ARCHITECTURE.md · docs/VALIDATION.md · docs/TEST_PLAN.md",
    },
    {
        "section": "KẾT LUẬN", "title": "Kết luận và hành động tiếp theo",
        "message": "Mô hình chứng minh được vòng giám sát end-to-end khi hoàn tất nghiệm thu collector, đường truyền và kênh nhận tin.",
        "points": [
            ("Chạy lab", "Xác nhận TAP/GNS3 và khôi phục Test-NetConnection TCP/22."),
            ("Nghiệm thu", "Chạy các ca ens33, CPU/RAM, mất SSH và phục hồi."),
            ("Báo cáo", "Ghi rõ SSH đã dùng; chỉ tuyên bố SNMP/Telegram khi đã đo thật."),
        ],
        "evidence": "THÀNH CÔNG  Cảnh báo đúng · phục hồi đúng · không spam · secret an toàn",
        "visual": {"kind": "actions"},
        "source": "docs/TEST_PLAN.md · docs/VALIDATION.md · docs/ARCHITECTURE.md",
    },
]


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, *, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line=None, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    sh.line.color.rgb = color(line or fill)
    sh.line.width = Pt(0.8)
    return sh


def text(slide, value, x, y, w, h, size=12, *, col=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.04,
         fit=True, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    if fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    return box


def rich_text(slide, term, detail, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    a = p.add_run()
    a.text = term + ": "
    a.font.name = FONT
    a.font.size = Pt(14.5)
    a.font.bold = True
    a.font.color.rgb = color(NAVY)
    b = p.add_run()
    b.text = detail
    b.font.name = FONT
    b.font.size = Pt(12.2)
    b.font.color.rgb = color(INK)
    return box


def connector(slide, x1, y1, x2, y2, col=MUTED, width=1.6):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color(col)
    line.line.width = Pt(width)
    try:
        from pptx.enum.dml import MSO_ARROWHEAD
        line.line.end_arrowhead = MSO_ARROWHEAD.TRIANGLE
    except (ImportError, AttributeError):
        pass
    return line


def node(slide, x, y, w, h, title, subtitle="", *, fill=PALE_BLUE,
         edge=BLUE, title_col=NAVY, title_size=13.0, subtitle_size=10.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    r = rect(slide, x, y, w, h, fill, shape=shape, line=edge)
    if subtitle:
        text(slide, title, x + 0.08, y + 0.08, w - 0.16, h * 0.40,
             title_size, col=title_col, bold=True, align=PP_ALIGN.CENTER, margin=0.01)
        text(slide, subtitle, x + 0.09, y + h * 0.49, w - 0.18, h * 0.38,
             subtitle_size, col=MUTED, align=PP_ALIGN.CENTER, margin=0.01)
    else:
        text(slide, title, x + 0.09, y + 0.05, w - 0.18, h - 0.10,
             title_size, col=title_col, bold=True, align=PP_ALIGN.CENTER, margin=0.01)
    return r


def panel_title(slide, x, y, label):
    text(slide, label.upper(), x + 0.18, y + 0.07, 5.5, 0.22, 9,
         col=TEAL_DARK, bold=True, margin=0)


def pill(slide, x, y, w, label, *, fill=PALE_TEAL, col=TEAL_DARK, size=9.5):
    rect(slide, x, y, w, 0.30, fill, line=fill)
    text(slide, label, x + 0.06, y + 0.015, w - 0.12, 0.27, size,
         col=col, bold=True, align=PP_ALIGN.CENTER, margin=0)


def draw_visual(slide, model, x, y, w, h):
    v = model["visual"]
    kind = v["kind"]
    if kind in {"cover", "flow"}:
        labels = v["nodes"]
        gap = 0.45
        nw = (w - (len(labels) - 1) * gap - 0.4) / len(labels)
        left = x + 0.2
        yy = y + 0.50
        hh = 0.82
        fills = [PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_TEAL]
        edges = [BLUE, ORANGE, GREEN, TEAL]
        for i, lab in enumerate(labels):
            xx = left + i * (nw + gap)
            node(slide, xx, yy, nw, hh, lab, fill=fills[i % 4], edge=edges[i % 4], title_size=11.4)
            if i < len(labels) - 1:
                connector(slide, xx + nw + 0.04, yy + hh / 2, xx + nw + gap - 0.04, yy + hh / 2)
        if kind == "cover":
            text(slide, "ĐỌC ĐỊNH KỲ", x + 0.28, y + 1.62, 2.0, 0.22, 9, col=BLUE, bold=True)
            text(slide, "ĐÁNH GIÁ", x + w * 0.36, y + 1.62, 1.6, 0.22, 9, col=ORANGE, bold=True)
            text(slide, "GỬI SỰ KIỆN", x + w * 0.72, y + 1.62, 2.0, 0.22, 9, col=TEAL_DARK, bold=True)
        return

    if kind == "architecture":
        row1 = [("Windows", "Python · monitor.py", PALE_BLUE, BLUE),
                ("GNS3", "Cloud · Switch1 · Cloud", PALE_ORANGE, ORANGE),
                ("Ubuntu VM", "CPU · RAM · interface", PALE_GREEN, GREEN)]
        xx = [x + 0.35, x + 4.55, x + 8.74]
        for i, (a, b, f, e) in enumerate(row1):
            node(slide, xx[i], y + 0.27, 3.35, 0.68, a, b, fill=f, edge=e)
        for i in range(2):
            connector(slide, xx[i] + 3.42, y + 0.62, xx[i + 1] - 0.10, y + 0.62)
        labels = [("Đánh giá ngưỡng", x + 0.65), ("SQLite / outbox", x + 4.85), ("Telegram / Slack", x + 9.03)]
        for i, (lab, px) in enumerate(labels):
            node(slide, px, y + 1.53, 2.95, 0.56, lab, fill=[PALE_ORANGE, PALE_PURPLE, PALE_TEAL][i], edge=[ORANGE, PURPLE, TEAL][i], title_size=11)
        connector(slide, x + 2.0, y + 0.98, x + 2.0, y + 1.36, ORANGE)
        connector(slide, x + 3.7, y + 1.80, x + 4.72, y + 1.80, PURPLE)
        connector(slide, x + 7.81, y + 1.80, x + 8.88, y + 1.80, TEAL)
        return

    if kind == "roles":
        cards = [
            ("01  MÁY THẬT", "Windows · VS Code\nPython + Internet", PALE_BLUE, BLUE),
            ("02  MẠNG LAB", "GNS3 Cloud + Switch\nVMware VMnet", PALE_ORANGE, ORANGE),
            ("03  MÁY ĐÍCH", "Ubuntu VM\nLinux /proc + /sys", PALE_GREEN, GREEN),
            ("04  NGƯỜI NHẬN", "Telegram / Slack\nỨng dụng điện thoại", PALE_TEAL, TEAL),
        ]
        for i, (a, b, f, e) in enumerate(cards):
            xx = x + 0.23 + i * 3.08
            node(slide, xx, y + 0.45, 2.72, 1.20, a, b, fill=f, edge=e, title_size=11.3, subtitle_size=9.8)
            if i < 3:
                connector(slide, xx + 2.76, y + 1.05, xx + 3.02, y + 1.05)
        text(slide, "Đo chỉ số: Ubuntu     •     Đánh giá/gửi tin: Windows     •     Chuyển tiếp: GNS3",
             x + 0.35, y + 2.04, w - 0.7, 0.26, 10, col=MUTED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "network":
        # Four clear columns: Windows adapter/helper, GNS3, Linux NICs, alert route.
        cols = [x + 0.15, x + 3.00, x + 6.00, x + 9.25]
        node(slide, cols[0], y + 0.33, 2.30, 0.58, "Windows TAP", "192.168.10.1/24", fill=PALE_BLUE, edge=BLUE, title_size=11)
        node(slide, cols[0], y + 1.10, 2.30, 0.58, "UDP helper", "127.0.0.1:30011", fill=PALE_BLUE, edge=BLUE, title_size=11)
        connector(slide, cols[0] + 1.15, y + 0.94, cols[0] + 1.15, y + 1.05, BLUE)
        node(slide, cols[1], y + 0.34, 2.55, 0.55, "Cloud-WINDOWS", "UDP local 30010", fill=PALE_ORANGE, edge=ORANGE, title_size=10.8)
        node(slide, cols[1], y + 1.07, 2.55, 0.55, "Switch1", "Hai port access VLAN 1", fill=PALE_ORANGE, edge=ORANGE, title_size=10.8)
        node(slide, cols[1], y + 1.80, 2.55, 0.55, "Cloud-LINUX", "VMware VMnet2", fill=PALE_ORANGE, edge=ORANGE, title_size=10.8)
        connector(slide, cols[1] + 1.275, y + 0.91, cols[1] + 1.275, y + 1.02, ORANGE)
        connector(slide, cols[1] + 1.275, y + 1.64, cols[1] + 1.275, y + 1.75, ORANGE)
        node(slide, cols[2], y + 0.40, 2.60, 0.78, "ens37 · QUẢN TRỊ", "192.168.10.129/24 · SSH 22", fill=PALE_GREEN, edge=GREEN, title_size=10.8)
        node(slide, cols[2], y + 1.51, 2.60, 0.78, "ens33 · THEO DÕI", "NAT .73.129/24* · VMnet8", fill=PALE_GREEN, edge=GREEN, title_size=10.8)
        connector(slide, cols[1] + 2.59, y + 2.07, cols[2] - 0.12, y + 0.79, GREEN)
        node(slide, cols[3], y + 0.87, 2.20, 0.84, "Internet Windows", "HTTPS → Telegram/Slack", fill=PALE_TEAL, edge=TEAL, title_size=10.5)
        connector(slide, cols[0] + 2.34, y + 0.61, cols[1] - 0.08, y + 0.61, BLUE)
        connector(slide, cols[0] + 2.34, y + 1.39, cols[1] - 0.08, y + 1.35, BLUE)
        text(slide, "* Địa chỉ DHCP đã ghi nhận; có thể thay đổi. TAP/helper là hạ tầng kết nối, không phải chức năng giám sát.",
             x + 0.2, y + 2.53, w - 0.4, 0.20, 8.7, col=MUTED, align=PP_ALIGN.CENTER)
        return

    if kind == "tappath":
        labels = [("Python", "Yêu cầu SSH", PALE_BLUE, BLUE),
                  ("TAP + UDP", "Chuyển frame", PALE_BLUE, BLUE),
                  ("Cloud + Switch", "GNS3", PALE_ORANGE, ORANGE),
                  ("ens37", "Ubuntu · TCP/22", PALE_GREEN, GREEN)]
        left = x + 0.25
        gap = 0.48
        nw = (w - 0.5 - 3 * gap) / 4
        for i, (a, b, f, e) in enumerate(labels):
            xx = left + i * (nw + gap)
            node(slide, xx, y + 0.48, nw, 0.90, a, b, fill=f, edge=e, title_size=11)
            if i < 3:
                connector(slide, xx + nw + 0.04, y + 0.93, xx + nw + gap - 0.04, y + 0.93)
        text(slide, "Lượt hỏi: Windows → Ubuntu      |      Kết quả: Ubuntu → Windows      |      HTTPS cảnh báo đi riêng qua Internet",
             x + 0.38, y + 1.75, w - 0.76, 0.35, 9.5, col=MUTED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "protocols":
        data = [("SSH", "TCP / 22", "Lệnh trên thiết bị", "Bản hiện đang chạy", PALE_BLUE, BLUE),
                ("SNMP", "UDP / 161", "Agent + OID", "Cần cài/cấu hình agent", PALE_ORANGE, ORANGE),
                ("HYBRID", "SSH + SNMP", "Nguồn ưu tiên/dự phòng", "Collector mở rộng", PALE_PURPLE, PURPLE)]
        gap = 0.30
        bw = (w - 0.65 - 2 * gap) / 3
        for i, (a, b, c, d, f, e) in enumerate(data):
            xx = x + 0.32 + i * (bw + gap)
            node(slide, xx, y + 0.31, bw, 1.27, a, b + "\n" + c, fill=f, edge=e, title_size=14, subtitle_size=9.4)
            pill(slide, xx + 0.28, y + 1.78, bw - 0.56, d, fill=f, col=e, size=8.8)
        return

    if kind == "modules":
        modules = [("monitor.py", "Vòng đo / giao diện", BLUE),
                   ("collectors", "SSH · SNMP · hybrid", GREEN),
                   ("engine.py", "Ngưỡng / trạng thái", ORANGE),
                   ("storage.py", "SQLite / outbox", PURPLE),
                   ("notifications.py", "Telegram / Slack", TEAL),
                   ("ui.py", "Bảng Rich Terminal", NAVY2)]
        bw, bh, gapx, gapy = 3.74, 0.75, 0.24, 0.22
        for i, (a, b, e) in enumerate(modules):
            row, col = divmod(i, 3)
            xx = x + 0.33 + col * (bw + gapx)
            yy = y + 0.31 + row * (bh + gapy)
            node(slide, xx, yy, bw, bh, a, b, fill=PALE_BLUE if row == 0 else PALE_TEAL, edge=e, title_size=11.4, subtitle_size=9.5)
        return

    if kind == "metrics":
        vals = [("CPU", "/proc/stat", "Bộ đếm bận", BLUE),
                ("RAM", "/proc/meminfo", "MemAvailable", PURPLE),
                ("TEMP", "/sys/hwmon", "UNKNOWN nếu thiếu", ORANGE),
                ("INTERFACE", "/sys/class/net", "operstate ens33", GREEN)]
        bw = (w - 0.7 - 3 * 0.18) / 4
        for i, (a, b, c, e) in enumerate(vals):
            xx = x + 0.35 + i * (bw + 0.18)
            node(slide, xx, y + 0.56, bw, 1.06, a, b + "\n" + c, fill=[PALE_BLUE, PALE_PURPLE, PALE_ORANGE, PALE_GREEN][i], edge=e, title_size=12, subtitle_size=9.3)
        return

    if kind == "formulas":
        rect(slide, x + 0.42, y + 0.25, w - 0.84, 0.53, NAVY, line=NAVY)
        text(slide, "CPU% = 100 × (1 − idle delta / total delta)", x + 0.65, y + 0.32, w - 1.3, 0.38, 15, col=WHITE, bold=True, align=PP_ALIGN.CENTER)
        rect(slide, x + 0.42, y + 0.93, w - 0.84, 0.53, PALE_PURPLE, line=PALE_PURPLE)
        text(slide, "RAM% = 100 × (MemTotal − MemAvailable) / MemTotal", x + 0.65, y + 1.00, w - 1.3, 0.38, 14, col=NAVY, bold=True, align=PP_ALIGN.CENTER)
        pill(slide, x + 1.10, y + 1.72, 3.3, "CPU: hai mẫu, cách nhau ~1 giây", fill=PALE_BLUE, col=BLUE)
        pill(slide, x + 7.40, y + 1.72, 3.55, "4 GB − 1 GB khả dụng ≈ 75% dùng", fill=PALE_GREEN, col=GREEN)
        return

    if kind == "temperature":
        node(slide, x + 0.28, y + 0.54, 3.35, 1.06, "AUTO", "Đọc cảm biến thật", fill=PALE_BLUE, edge=BLUE, title_size=13)
        node(slide, x + 4.82, y + 0.54, 3.35, 1.06, "KHÔNG CÓ SENSOR", "UNKNOWN · giữ đúng dữ liệu", fill=PALE_ORANGE, edge=ORANGE, title_size=12)
        node(slide, x + 9.35, y + 0.54, 2.60, 1.06, "DEMO", "MÔ PHỎNG", fill=PALE_RED, edge=RED, title_size=13)
        connector(slide, x + 3.70, y + 1.07, x + 4.70, y + 1.07, ORANGE)
        text(slide, "Mô phỏng không chứng minh đã đo nhiệt độ phần cứng", x + 7.9, y + 1.78, 3.9, 0.26, 9, col=RED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "cycle":
        labels = [("1", "Đọc config", BLUE), ("2", "Kết nối SSH", BLUE), ("3", "Thu thập", GREEN),
                  ("4", "So ngưỡng", ORANGE), ("5", "Lưu state", PURPLE), ("6", "Gửi / chờ", TEAL)]
        gap = 0.18
        bw = (w - 0.58 - 5 * gap) / 6
        for i, (n, lab, e) in enumerate(labels):
            xx = x + 0.29 + i * (bw + gap)
            rect(slide, xx, y + 0.47, bw, 0.94, PALE_BLUE if i < 2 else PALE_GREEN if i == 2 else PALE_ORANGE if i == 3 else PALE_PURPLE if i == 4 else PALE_TEAL, line=e)
            pill(slide, xx + 0.12, y + 0.54, 0.35, n, fill=e, col=WHITE, size=9)
            text(slide, lab, xx + 0.07, y + 0.96, bw - 0.14, 0.32, 9.5, col=NAVY, bold=True, align=PP_ALIGN.CENTER)
            if i < 5:
                connector(slide, xx + bw + 0.02, y + 0.94, xx + bw + gap - 0.02, y + 0.94)
        text(slide, "5s sau lượt đọc  +  ~1s lấy mẫu CPU  +  thời gian SSH / timeout / API  →  độ trễ thay đổi",
             x + 0.35, y + 1.75, w - 0.70, 0.32, 10, col=MUTED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "thresholds":
        rows = [("CPU", 75, 70, BLUE), ("RAM", 65, 60, PURPLE), ("TEMP", 65, 60, ORANGE)]
        for i, (lab, alarm, rec, e) in enumerate(rows):
            yy = y + 0.28 + i * 0.59
            text(slide, lab, x + 0.40, yy, 0.90, 0.34, 11, col=NAVY, bold=True)
            bx = x + 1.45
            rect(slide, bx, yy + 0.10, 6.3, 0.12, LINE, shape=MSO_SHAPE.RECTANGLE, line=LINE)
            rect(slide, bx, yy + 0.10, 6.3 * alarm / 100, 0.12, e, shape=MSO_SHAPE.RECTANGLE, line=e)
            pill(slide, x + 8.05, yy + 0.01, 1.50, f">{alarm}{'°C' if lab=='TEMP' else '%'}", fill=PALE_RED, col=RED)
            pill(slide, x + 9.85, yy + 0.01, 2.15, f"≤{rec}{'°C' if lab=='TEMP' else '%'} phục hồi", fill=PALE_GREEN, col=GREEN)
        text(slide, "Interface: DOWN → ALERT     |     UP → phục hồi", x + 2.4, y + 2.12, 7.3, 0.28, 10, col=INK, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "states":
        seq = [("OK", "50%", PALE_GREEN, GREEN), ("ALERT", "72%", PALE_RED, RED), ("ALERT", "74%", PALE_RED, RED),
               ("ALERT", "63%", PALE_ORANGE, ORANGE), ("OK", "58%", PALE_GREEN, GREEN)]
        bw, gap = 1.73, 0.63
        start = x + 0.25
        for i, (s, val, f, e) in enumerate(seq):
            xx = start + i * (bw + gap)
            node(slide, xx, y + 0.42, bw, 0.78, s, f"RAM {val}", fill=f, edge=e, title_size=11.2, subtitle_size=9.5)
            if i < len(seq) - 1:
                connector(slide, xx + bw + 0.03, y + 0.81, xx + bw + gap - 0.03, y + 0.81)
        text(slide, "Chỉ gửi khi chuyển trạng thái: vượt 65% tạo alert; ≤60% mới phục hồi; 63% còn ALERT nếu trước đó đã cảnh báo.",
             x + 0.38, y + 1.57, w - 0.76, 0.40, 9.7, col=MUTED, bold=True, align=PP_ALIGN.CENTER)
        node(slide, x + 9.78, y + 2.08, 2.25, 0.43, "UNKNOWN", "thiếu dữ liệu", fill=PALE_BLUE, edge=BLUE, title_size=9.6, subtitle_size=8.2)
        return

    if kind == "incidents":
        rows = [("ens33 DOWN", "ens37 SSH còn", "Cảnh báo interface", PALE_ORANGE, ORANGE),
                ("Mất SSH/GNS3", "Không có mẫu mới", "Cảnh báo kết nối", PALE_BLUE, BLUE),
                ("Mất Internet", "Tin lưu trong outbox", "Gửi lại khi có mạng", PALE_TEAL, TEAL)]
        for i, (a, b, c, f, e) in enumerate(rows):
            yy = y + 0.28 + i * 0.63
            node(slide, x + 0.50, yy, 2.85, 0.48, a, fill=f, edge=e, title_size=10.3)
            connector(slide, x + 3.43, yy + 0.24, x + 4.03, yy + 0.24, MUTED)
            node(slide, x + 4.12, yy, 3.50, 0.48, b, fill=PALE_BLUE, edge=BLUE, title_size=10.0)
            connector(slide, x + 7.72, yy + 0.24, x + 8.32, yy + 0.24, MUTED)
            node(slide, x + 8.41, yy, 3.35, 0.48, c, fill=PALE_GREEN, edge=GREEN, title_size=10.0)
        return

    if kind == "outbox":
        labels = [("STATE", "Chuyển trạng thái", BLUE), ("SQLITE", "Commit + outbox", PURPLE),
                  ("SENDER", "Luồng riêng", TEAL), ("API", "Xác nhận", GREEN)]
        gap = 0.55
        bw = (w - 0.7 - gap * 3) / 4
        for i, (a, b, e) in enumerate(labels):
            xx = x + 0.35 + i * (bw + gap)
            node(slide, xx, y + 0.39, bw, 0.90, a, b, fill=[PALE_BLUE, PALE_PURPLE, PALE_TEAL, PALE_GREEN][i], edge=e, title_size=11.2)
            if i < 3:
                connector(slide, xx + bw + 0.03, y + 0.84, xx + bw + gap - 0.03, y + 0.84)
        node(slide, x + 4.25, y + 1.66, 4.0, 0.43, "HTTP lỗi / 429", "giữ tin · backoff · thử lại", fill=PALE_ORANGE, edge=ORANGE, title_size=9.6, subtitle_size=8.3)
        connector(slide, x + 8.45, y + 1.90, x + 10.2, y + 1.31, ORANGE)
        text(slide, "At-least-once: xác nhận bị mất có thể gây gửi trùng", x + 8.45, y + 1.69, 3.35, 0.34, 8.7, col=RED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "security":
        cards = [("01", "Host key", "Xác minh fingerprint", BLUE, PALE_BLUE),
                 ("02", "Secret", ".env cục bộ · Git ignore", PURPLE, PALE_PURPLE),
                 ("03", "Kênh", "HTTPS · TLS · không log URL", TEAL, PALE_TEAL),
                 ("04", "SNMP", "v3 authPriv / v2c lab cô lập", ORANGE, PALE_ORANGE)]
        bw = (w - 0.74 - 3 * 0.18) / 4
        for i, (n, a, b, e, f) in enumerate(cards):
            xx = x + 0.37 + i * (bw + 0.18)
            node(slide, xx, y + 0.54, bw, 1.07, a, b, fill=f, edge=e, title_size=11.3, subtitle_size=9.0)
            pill(slide, xx + 0.45, y + 0.30, 0.48, n, fill=e, col=WHITE, size=8.8)
        return

    if kind == "profiles":
        node(slide, x + 0.35, y + 0.32, 4.65, 1.67, "BẢN ĐƠN GIẢN", "SSH · .env.simple · ens33\nWindows → .10.129\nrun_windows.cmd / F5",
             fill=PALE_BLUE, edge=BLUE, title_size=13, subtitle_size=10.2)
        node(slide, x + 6.65, y + 0.32, 4.65, 1.67, "PROFILE GNS3 MỞ RỘNG", "SSH / SNMP / hybrid · .env.gns3\ngns3test .30.129 · VPCS .30.10\nrun: python -m device_monitor",
             fill=PALE_ORANGE, edge=ORANGE, title_size=12.2, subtitle_size=10)
        pill(slide, x + 5.15, y + 0.67, 1.35, "CHUNG", fill=PALE_TEAL, col=TEAL_DARK, size=9.5)
        text(slide, "engine · SQLite/outbox · dashboard · Telegram/Slack", x + 2.62, y + 2.17, 6.4, 0.27, 9.5, col=MUTED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "testmatrix":
        headers = ["Kích thích", "Kết quả cần thấy", "Bằng chứng"]
        widths = [3.6, 4.7, 3.5]
        xs = [x + 0.3, x + 3.97, x + 8.74]
        for xx, ww, label in zip(xs, widths, headers):
            rect(slide, xx, y + 0.22, ww, 0.32, NAVY, shape=MSO_SHAPE.RECTANGLE, line=NAVY)
            text(slide, label, xx + .1, y + .24, ww - .2, .26, 9, col=WHITE, bold=True)
        rows = [("CPU/RAM tăng", "ALERT → phục hồi", "Số đo + tin"),
                ("ens33 / SSH ngắt", "DOWN / mất kết nối", "State + thời điểm"),
                ("Internet / restart", "Queue + retry", "Tin chờ → giao"),
                ("Nhiệt độ / host key", "UNKNOWN / từ chối", "Nguồn + fingerprint")]
        for i, row in enumerate(rows):
            yy = y + 0.57 + i * 0.38
            bg = WHITE if i % 2 == 0 else BG
            for j, (xx, ww, val) in enumerate(zip(xs, widths, row)):
                rect(slide, xx, yy, ww, 0.34, bg, shape=MSO_SHAPE.RECTANGLE, line=LINE)
                text(slide, val, xx + 0.10, yy + 0.015, ww - .18, .29, 8.7, col=INK, bold=(j == 0))
        return

    if kind == "evidence":
        stats = [("TCP/22", "PASS", GREEN, PALE_GREEN),
                 ("SOURCE", ".10.1 / TAP", BLUE, PALE_BLUE),
                 ("106 / 57", "frame counters", TEAL, PALE_TEAL),
                 ("DROPPED", "0 (ảnh chụp)", ORANGE, PALE_ORANGE)]
        bw = (w - 0.78 - 3 * 0.18) / 4
        for i, (a, b, e, f) in enumerate(stats):
            xx = x + 0.39 + i * (bw + 0.18)
            node(slide, xx, y + 0.48, bw, 1.05, a, b, fill=f, edge=e, title_size=13.0, subtitle_size=9.2)
        text(slide, "Một lần thử mạng thành công ≠ nghiệm thu toàn bộ collector / SNMP / gửi tin thật",
             x + 0.52, y + 1.81, w - 1.04, 0.30, 9.2, col=RED, bold=True, align=PP_ALIGN.CENTER)
        return

    if kind == "limitations":
        vals = [("CẢM BIẾN", "VM có thể thiếu nhiệt độ", ORANGE, PALE_ORANGE),
                ("HISTORIAN", "SQLite không lưu chuỗi thời gian", PURPLE, PALE_PURPLE),
                ("SINGLE HOST", "Host ngủ/tắt: monitor dừng", RED, PALE_RED),
                ("SCOPE", "Một tiến trình mỗi state file", BLUE, PALE_BLUE)]
        for i, (a, b, e, f) in enumerate(vals):
            row, col = divmod(i, 2)
            xx = x + 0.65 + col * 6.13
            yy = y + 0.30 + row * 0.97
            node(slide, xx, yy, 5.48, 0.72, a, b, fill=f, edge=e, title_size=10.3, subtitle_size=9.2)
        return

    if kind == "actions":
        items = [("1", "Khôi phục TAP/GNS3", BLUE), ("2", "Lặp thử TCP/22", GREEN),
                 ("3", "Chạy collector / --once", ORANGE), ("4", "Nhận cảnh báo thật", TEAL),
                 ("5", "Ghi log nghiệm thu", PURPLE)]
        gap = 0.31
        bw = (w - 0.66 - 4 * gap) / 5
        for i, (n, lab, e) in enumerate(items):
            xx = x + 0.33 + i * (bw + gap)
            rect(slide, xx, y + 0.48, bw, 1.08, WHITE, line=e)
            pill(slide, xx + bw / 2 - 0.19, y + 0.34, 0.38, n, fill=e, col=WHITE, size=9)
            text(slide, lab, xx + 0.09, y + 0.76, bw - 0.18, 0.55, 9.5, col=NAVY, bold=True, align=PP_ALIGN.CENTER)
            if i < 4:
                connector(slide, xx + bw + 0.02, y + 1.02, xx + bw + gap - 0.02, y + 1.02)
        return

    raise ValueError(f"Unknown visual kind: {kind}")


def add_cover(slide, model):
    rect(slide, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE, line=NAVY)
    rect(slide, 0, 0, 0.16, H, TEAL, shape=MSO_SHAPE.RECTANGLE, line=TEAL)
    text(slide, "QTMNC  /  EXECUTIVE BRIEFING", 0.78, 0.48, 6.2, 0.34, 11, col="99E7DA", bold=True)
    text(slide, model["title"], 0.78, 1.07, 11.8, 1.16, 34, col=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    text(slide, "Từ dữ liệu thiết bị đến cảnh báo có thể hành động",
         0.82, 2.30, 11.6, 0.45, 18, col="CCD6E2")
    rect(slide, 0.80, 3.03, 11.74, 0.80, NAVY2, line="305271")
    text(slide, "THÔNG ĐIỆP CỐT LÕI", 1.05, 3.12, 2.4, 0.20, 9, col="99E7DA", bold=True)
    text(slide, model["message"], 1.05, 3.36, 11.13, 0.34, 13.8, col=WHITE, bold=True)
    points = model["points"]
    for i, (head, body) in enumerate(points):
        xx = 0.82 + i * 4.0
        rect(slide, xx, 4.25, 3.72, 0.85, NAVY2, line="305271")
        text(slide, head.upper(), xx + .16, 4.33, 3.36, .22, 9, col="99E7DA", bold=True)
        text(slide, body, xx + .16, 4.57, 3.36, .40, 10.0, col=WHITE)
    v = model["visual"]
    labels = v["nodes"]
    gap = 0.48
    nw = (11.7 - (len(labels) - 1) * gap) / len(labels)
    left, yy, hh = .84, 5.73, .77
    for i, label in enumerate(labels):
        xx = left + i * (nw + gap)
        c = [BLUE, ORANGE, GREEN, TEAL][i]
        fill = ["223758", "3D342B", "243D34", "203E40"][i]
        node(slide, xx, yy, nw, hh, label, fill=fill, edge=c, title_col=WHITE, title_size=10.7)
        if i < len(labels) - 1:
            connector(slide, xx + nw + .04, yy + hh / 2, xx + nw + gap - .04, yy + hh / 2, "89A0B7")
    text(slide, "BẢN ĐƠN GIẢN ĐANG DÙNG SSH  ·  SNMP/HYBRID LÀ PROFILE MỞ RỘNG",
         .86, 6.86, 11.3, .25, 9, col="AFC0D0", bold=True)
    text(slide, "01 / 22", 11.65, 7.10, 0.80, .20, 9, col="AFC0D0", align=PP_ALIGN.RIGHT)


def add_regular_slide(prs, model, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    rect(slide, 0, 0, W, 0.12, TEAL, shape=MSO_SHAPE.RECTANGLE, line=TEAL)
    text(slide, f"QTMNC   /   {model['section']}", .65, .28, 8.5, .24, 9.2, col=TEAL_DARK, bold=True, margin=0)
    text(slide, model["title"], .63, .59, 11.9, .57, 27, col=NAVY, bold=True, margin=0)
    pill(slide, 12.15, .42, .55, f"{idx:02d}", fill=NAVY, col=WHITE, size=9)

    rect(slide, .65, 1.28, 12.04, .57, WHITE, line=LINE)
    text(slide, "THÔNG ĐIỆP", .83, 1.33, 1.36, .15, 8, col=TEAL_DARK, bold=True, margin=0)
    text(slide, model["message"], 2.15, 1.34, 10.3, .42, 13.0, col=NAVY, bold=True, margin=.01)

    card_y, card_h = 2.02, 1.13
    colors = [BLUE, ORANGE, TEAL]
    fills = [PALE_BLUE, PALE_ORANGE, PALE_TEAL]
    gap = .20
    card_w = (12.04 - 2 * gap) / 3
    for i, (head, body) in enumerate(model["points"]):
        xx = .65 + i * (card_w + gap)
        rect(slide, xx, card_y, card_w, card_h, fills[i], line=fills[i])
        rect(slide, xx, card_y, .055, card_h, colors[i], shape=MSO_SHAPE.RECTANGLE, line=colors[i])
        text(slide, f"{i + 1:02d}  {head}", xx + .20, card_y + .10, card_w - .37, .28, 12.4, col=NAVY, bold=True, margin=0)
        text(slide, body, xx + .20, card_y + .42, card_w - .38, .59, 10.4, col=INK, valign=MSO_ANCHOR.TOP, margin=.015)

    rect(slide, .65, 3.27, 12.04, .39, NAVY, line=NAVY)
    text(slide, "MINH CHỨNG", .82, 3.31, 1.20, .20, 8.0, col="9CE2D6", bold=True, margin=0)
    text(slide, model["evidence"], 2.05, 3.31, 10.4, .23, 9.4, col=WHITE, bold=True, margin=0)

    rect(slide, .65, 3.80, 12.04, 2.99, WHITE, line=LINE)
    panel_title(slide, .65, 3.80, "Gợi ý trực quan  /  " + model["visual"]["kind"].upper())
    draw_visual(slide, model, .84, 4.13, 11.65, 2.50)

    text(slide, "NGUỒN  " + model["source"], .68, 6.91, 11.6, .20, 7.5, col=MUTED, margin=0)
    text(slide, f"{idx:02d} / 22", 11.85, 7.12, .82, .20, 8.5, col=MUTED, align=PP_ALIGN.RIGHT, margin=0)
    return slide


def build_deck():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Hệ thống giám sát thiết bị và cảnh báo Telegram/Slack"
    prs.core_properties.subject = "Dàn ý thuyết trình Executive 22 slide · Windows/GNS3/Ubuntu"
    prs.core_properties.author = "QTMNC"
    prs.core_properties.keywords = "GNS3, SSH, SNMP, Telegram, Slack, Ubuntu"
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover_20(cover, SLIDES[0])
    for i, model in enumerate(SLIDES[1:], start=2):
        add_regular_slide_20(prs, model, i)
    out = OUT / "QTMNC_Executive_Thuyet_trinh.pptx"
    prs.save(out)
    return out


def build_outline():
    lines = [
        "# Dàn ý thuyết trình: Hệ thống giám sát thiết bị",
        "",
        "**Bộ trình chiếu:** 22 slide · 16:9 · cỡ chữ trên slide tối thiểu 20 pt · phần tử PowerPoint có thể chỉnh sửa.",
        "**Phạm vi trình bày:** Windows chạy Python → GNS3 → Ubuntu VMware → Telegram/Slack.",
        "**Lưu ý trạng thái:** Bản đơn giản đã cấu hình SSH; SNMP/hybrid là phương án mở rộng. Kết quả loopback tự động không thay thế nghiệm thu agent/API trên thiết bị thật.",
        "",
    ]
    for i, model in enumerate(SLIDES, start=1):
        lines.extend([f"## SLIDE {i}: {model['title']}", "",
                      f"- **Thông điệp cốt lõi:** {model['message']}"])
        for n, (head, body) in enumerate(model["points"], start=1):
            lines.append(f"- **Điểm chính {n}: {head}** — {body}")
        lines.extend([f"- **Số liệu/Minh chứng:** {model['evidence']}",
                      f"- **Gợi ý trực quan:** {visual_name(model['visual']['kind'])}.",
                      f"- **Nguồn:** {model['source']}", ""])
    path = OUT / "QTMNC_Dan_y_22_slide.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def visual_name(kind: str) -> str:
    return {
        "cover": "Sơ đồ luồng điều hành 4 tác nhân",
        "flow": "Sơ đồ người kiểm tra – hệ thống – người nhận",
        "architecture": "Kiến trúc end-to-end dạng 2 tầng",
        "roles": "Thẻ vai trò 4 khối thiết bị",
        "network": "Sơ đồ mạng phân cột Windows – GNS3 – Ubuntu – Internet",
        "tappath": "Luồng yêu cầu và chiều trả lời SSH",
        "protocols": "So sánh 3 giao thức bằng thẻ",
        "modules": "Sơ đồ module dạng lưới",
        "metrics": "Bốn thẻ metric và nguồn dữ liệu",
        "formulas": "Hai công thức và ví dụ định lượng",
        "temperature": "Ba nhánh auto – thiếu sensor – demo",
        "cycle": "Timeline 6 bước trong vòng đo",
        "thresholds": "Thanh ngưỡng 4 metric",
        "states": "Dòng trạng thái RAM với hysteresis",
        "incidents": "Ma trận 3 loại sự cố và phản ứng",
        "outbox": "Luồng state → SQLite → sender → API/retry",
        "security": "Bốn lớp xác minh và bảo vệ",
        "profiles": "So sánh hai profile đơn giản/mở rộng",
        "testmatrix": "Ma trận kích thích – kết quả – bằng chứng",
        "evidence": "Thẻ số liệu kiểm chứng thực tế",
        "limitations": "Bốn khối giới hạn kỹ thuật",
        "actions": "Lộ trình 5 bước nghiệm thu",
    }[kind]


# Large-type presentation layout. All visible text on these slides is >= 20 pt.
def large_text(slide, value, x, y, w, h, *, size=20, col=INK, bold=False,
               align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.06):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.015)
    tf.margin_bottom = Inches(0.015)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(max(20, size))
    run.font.bold = bold
    run.font.color.rgb = color(col)
    return box


def large_node(slide, x, y, w, h, heading, detail="", *, fill=PALE_BLUE,
               edge=BLUE, dark=False):
    rect(slide, x, y, w, h, fill, line=edge)
    main_col = WHITE if dark else NAVY
    sub_col = "DFE8F1" if dark else MUTED
    if detail:
        large_text(slide, heading, x + .08, y + .08, w - .16, .43,
                   col=main_col, bold=True, align=PP_ALIGN.CENTER, margin=.02)
        large_text(slide, detail, x + .08, y + .53, w - .16, h - .60,
                   col=sub_col, align=PP_ALIGN.CENTER, margin=.02)
    else:
        large_text(slide, heading, x + .08, y + .07, w - .16, h - .14,
                   col=main_col, bold=True, align=PP_ALIGN.CENTER, margin=.02)


def large_flow(slide, x, y, w, h, items, *, palette=None, dark=False):
    palette = palette or [BLUE, ORANGE, GREEN, TEAL]
    gap = .35 if len(items) <= 3 else .25
    node_w = (w - (len(items) - 1) * gap) / len(items)
    node_h = min(h, 1.10)
    yy = y + (h - node_h) / 2
    for i, (heading, detail) in enumerate(items):
        xx = x + i * (node_w + gap)
        col = palette[i % len(palette)]
        fill = ("223758", "3D342B", "243D34", "203E40")[i % 4] if dark else (PALE_BLUE, PALE_ORANGE, PALE_GREEN, PALE_TEAL)[i % 4]
        large_node(slide, xx, yy, node_w, node_h, heading, detail,
                   fill=fill, edge=col, dark=dark)
        if i < len(items) - 1:
            connector(slide, xx + node_w + .04, yy + node_h / 2,
                      xx + node_w + gap - .04, yy + node_h / 2, col, width=2.2)


def large_visual(slide, model, x, y, w, h):
    kind = model["visual"]["kind"]
    visuals = {
        "cover": [("WINDOWS", "PYTHON / VS CODE"), ("GNS3", "MẠNG LAB"), ("UBUNTU", "THIẾT BỊ ĐÍCH"), ("BOT API", "TELEGRAM / SLACK")],
        "flow": [("THIẾT BỊ", "ĐỌC CHỈ SỐ"), ("PYTHON", "ĐÁNH GIÁ"), ("QUẢN TRỊ", "NHẬN TIN")],
        "architecture": [("WINDOWS", "PYTHON"), ("GNS3", "SSH / TCP 22"), ("UBUNTU", "CPU · RAM · NIC"), ("BOT API", "HTTPS · ALERTS")],
        "roles": [("WINDOWS", "GIÁM SÁT"), ("GNS3", "CHUYỂN TIẾP"), ("UBUNTU", "ĐỐI TƯỢNG ĐO"), ("BOT", "GỬI CẢNH BÁO")],
        "network": [("WINDOWS", "TAP · .10.1"), ("CLOUD", "GNS3"), ("SWITCH", "VLAN 1"), ("UBUNTU", "ens37 · .10.129")],
        "tappath": [("PYTHON", "YÊU CẦU SSH"), ("TAP / UDP", "FRAME"), ("GNS3", "SWITCH"), ("UBUNTU", "SSHD · 22")],
        "protocols": [("SSH", "TCP / 22"), ("SNMP", "UDP / 161"), ("HYBRID", "DỰ PHÒNG")],
        "modules": [("COLLECTOR", "SSH / SNMP"), ("ENGINE", "NGƯỠNG"), ("SQLITE", "STATE / OUTBOX"), ("SENDER", "TELEGRAM / SLACK")],
        "metrics": [("CPU", "/proc/stat"), ("RAM", "/proc/meminfo"), ("TEMP", "/sys/hwmon"), ("INTERFACE", "ens33")],
        "formulas": [("CPU", "100 × (1 − idle Δ / total Δ)"), ("RAM", "100 × (total − available) / total")],
        "temperature": [("AUTO", "ĐỌC CẢM BIẾN"), ("UNKNOWN", "THIẾU SENSOR"), ("DEMO", "NHÃN MÔ PHỎNG")],
        "cycle": [("CONNECT", "KẾT NỐI"), ("READ", "THU THẬP"), ("EVALUATE", "ĐÁNH GIÁ"), ("NOTIFY", "GỬI TIN")],
        "thresholds": [("CPU", "> 75%"), ("RAM", "> 65%"), ("TEMP", "> 65°C"), ("IFACE", "DOWN")],
        "states": [("OK", "BÌNH THƯỜNG"), ("ALERT", "VƯỢT NGƯỠNG"), ("RECOVERY", "ĐÃ PHỤC HỒI")],
        "incidents": [("ens33 DOWN", "SSH CÒN"), ("SSH MẤT", "UNKNOWN"), ("MẠNG MẤT", "QUEUE")],
        "outbox": [("STATE", "SỰ KIỆN"), ("SQLITE", "LƯU BỀN"), ("OUTBOX", "CHỜ GỬI"), ("API", "THỬ LẠI")],
        "security": [("HOST KEY", "XÁC MINH"), ("SECRET", ".ENV CỤC BỘ"), ("HTTPS", "TLS"), ("SNMPv3", "AUTH / PRIV")],
        "profiles": [("ĐƠN GIẢN", "SSH"), ("MỞ RỘNG", "SSH + SNMP")],
        "testmatrix": [("CPU / RAM", "VƯỢT NGƯỠNG"), ("SSH / NIC", "NGẮT KẾT NỐI"), ("INTERNET", "QUEUE / RETRY")],
        "evidence": [("119", "UNIT TESTS"), ("TCP / 22", "ĐÃ KẾT NỐI"), ("106 / 57", "FRAME"), ("0", "DROPPED")],
        "limitations": [("NHIỆT ĐỘ", "VM THIẾU SENSOR"), ("LỊCH SỬ", "SQLITE / STATE"), ("MÁY CHỦ", "HOST TẮT → DỪNG")],
        "actions": [("1 · MẠNG", "KHÔI PHỤC"), ("2 · SSH", "XÁC NHẬN"), ("3 · COLLECTOR", "CHẠY THẬT"), ("4 · ALERT", "NGHIỆM THU")],
    }
    large_flow(slide, x + .25, y, w - .50, h, visuals[kind])


def add_cover_20(slide, model):
    rect(slide, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE, line=NAVY)
    rect(slide, 0, 0, .16, H, TEAL, shape=MSO_SHAPE.RECTANGLE, line=TEAL)
    large_text(slide, "QTMNC  ·  THUYẾT TRÌNH DỰ ÁN", .78, .38, 8.0, .42,
               size=20, col="99E7DA", bold=True)
    large_text(slide, model["title"], .78, .90, 11.9, .90,
               size=38, col=WHITE, bold=True)
    large_text(slide, model["message"], .82, 1.89, 11.5, .82,
               size=22, col="E8EEF5")
    card_w, card_h, gap = 3.78, 1.48, .24
    for i, (head, body) in enumerate(model["points"]):
        xx = .82 + i * (card_w + gap)
        rect(slide, xx, 2.95, card_w, card_h, NAVY2, line="305271")
        large_text(slide, head.upper(), xx + .13, 3.04, card_w - .26, .34,
                   size=20, col="99E7DA", bold=True)
        large_text(slide, body, xx + .13, 3.40, card_w - .26, .90,
                   size=20, col=WHITE)
    large_flow(slide, .86, 4.73, 11.6, 1.10, [
        ("WINDOWS", "PYTHON / VS CODE"), ("GNS3", "MẠNG LAB"),
        ("UBUNTU", "THIẾT BỊ ĐÍCH"), ("BOT API", "TELEGRAM / SLACK")], dark=True)
    large_text(slide, model["evidence"], .86, 6.05, 11.55, .64,
               size=20, col="DFE8F1", bold=True, align=PP_ALIGN.CENTER)
    large_text(slide, "01 / 22", 10.70, 7.02, 1.80, .32,
               size=20, col="AFC0D0", align=PP_ALIGN.RIGHT)


def add_regular_slide_20(prs, model, idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    rect(slide, 0, 0, W, .13, TEAL, shape=MSO_SHAPE.RECTANGLE, line=TEAL)
    large_text(slide, f"QTMNC  /  {model['section']}", .64, .25, 8.8, .38,
               size=20, col=TEAL_DARK, bold=True, margin=.01)
    large_text(slide, model["title"], .63, .68, 11.0, .66,
               size=30, col=NAVY, bold=True, margin=.01)
    rect(slide, 11.85, .66, .84, .55, NAVY, line=NAVY)
    large_text(slide, f"{idx:02d}", 11.88, .70, .78, .46,
               size=20, col=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=.01)

    rect(slide, .65, 1.48, 12.04, .82, WHITE, line=LINE)
    large_text(slide, "THÔNG ĐIỆP  ·  " + model["message"], .82, 1.55, 11.68, .68,
               size=20, col=NAVY, bold=True, margin=.04)

    fills = [PALE_BLUE, PALE_ORANGE, PALE_TEAL]
    edges = [BLUE, ORANGE, TEAL]
    for i, (head, body) in enumerate(model["points"]):
        yy = 2.43 + i * .78
        rect(slide, .65, yy, 12.04, .68, fills[i], line=fills[i])
        rect(slide, .65, yy, .08, .68, edges[i], shape=MSO_SHAPE.RECTANGLE, line=edges[i])
        large_text(slide, f"{i + 1:02d}  {head}: {body}", .83, yy + .04, 11.70, .59,
                   size=20, col=NAVY, bold=(i == 0), margin=.025)

    rect(slide, .65, 4.88, 12.04, 1.72, WHITE, line=LINE)
    visual_titles = {"network": "SƠ ĐỒ MẠNG GNS3", "architecture": "KIẾN TRÚC TỔNG THỂ",
                     "cover": "LUỒNG ĐIỀU HÀNH", "testmatrix": "MA TRẬN KIỂM THỬ"}
    visual_title = visual_titles.get(model["visual"]["kind"], visual_name(model["visual"]["kind"]))
    large_text(slide, "GỢI Ý TRỰC QUAN  ·  " + visual_title.upper(),
               .84, 4.94, 11.60, .33, size=20, col=TEAL_DARK, bold=True, margin=.01)
    large_visual(slide, model, .86, 5.30, 11.62, 1.20)

    rect(slide, .65, 6.74, 12.04, .53, NAVY, line=NAVY)
    large_text(slide, model["evidence"], .82, 6.77, 11.68, .46,
               size=20, col=WHITE, bold=True, margin=.025)
    return slide


if __name__ == "__main__":
    deck = build_deck()
    outline = build_outline()
    print(deck)
    print(outline)
